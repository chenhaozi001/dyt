#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成「通话复盘助手」专题分享 PPT（只讲这一个工具）。

依赖：python-pptx, Pillow（生成示例截图）
  pip install python-pptx Pillow
  python tools/gen_screenshots.py   # 先生成截图
  python tools/build_pptx.py        # 再生成 PPT
"""

import argparse
import os

from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PRIMARY = RGBColor(0x1F, 0x4E, 0x79)
ACCENT = RGBColor(0x2E, 0x86, 0xC1)
DARK = RGBColor(0x22, 0x29, 0x2F)
GRAY = RGBColor(0x55, 0x5F, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

ROOT = os.path.join(os.path.dirname(__file__), "..")
SHOT_DIR = os.path.join(ROOT, "assets", "screenshots")


def _notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def _title_bar(slide, prs, title):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), prs.slide_width - Inches(1.2), Inches(0.9))
    tf = box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = PRIMARY
    r.font.name = "Microsoft YaHei"
    line = slide.shapes.add_shape(1, Inches(0.65), Inches(1.2), Inches(2.0), Pt(3))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def add_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_shape(1, 0, Inches(2.0), prs.slide_width, Inches(3.5)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = PRIMARY
    slide.shapes[-1].line.fill.background()

    t = slide.shapes.add_textbox(Inches(0.8), Inches(2.3), prs.slide_width - Inches(1.6), Inches(1.2))
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "AI 学习成果分享"
    r.font.size = Pt(40)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = "Microsoft YaHei"

    s = slide.shapes.add_textbox(Inches(0.8), Inches(3.5), prs.slide_width - Inches(1.6), Inches(1.0))
    for i, line in enumerate(["通话复盘助手", "—— 用 Cursor 把复盘从 20 分钟缩短到 2 分钟", "分享人：XXX"]):
        para = s.text_frame.paragraphs[0] if i == 0 else s.text_frame.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        rr = para.add_run()
        rr.text = line
        rr.font.size = Pt(22 if i == 0 else 16)
        rr.font.color.rgb = RGBColor(0xE8, 0xF1, 0xFB) if i else WHITE
        rr.font.bold = i == 0
        rr.font.name = "Microsoft YaHei"
    _notes(slide, "大家好，我是做电话销售的。这次分享我做的一个工具——通话复盘助手。"
                 "它帮我把每通电话的复盘时间从二十分钟缩短到两分钟。全程用 Cursor 开发，部署成了一个网站，打开就能用。")


def add_bullets(prs, title, bullets, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _title_bar(slide, prs, title)
    body = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), prs.slide_width - Inches(1.6), Inches(5.5))
    tf = body.text_frame
    tf.word_wrap = True
    for i, (text, level) in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.level = level
        para.space_after = Pt(8)
        r = para.add_run()
        r.text = ("• " if level == 0 else "– ") + text
        r.font.size = Pt(22 if level == 0 else 17)
        r.font.bold = level == 0
        r.font.color.rgb = DARK if level == 0 else GRAY
        r.font.name = "Microsoft YaHei"
    _notes(slide, notes)


def add_image_slide(prs, title, img_path, caption, notes, img_left=Inches(1.2), img_top=Inches(1.4), img_w=Inches(10.5)):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _title_bar(slide, prs, title)
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, img_left, img_top, width=img_w)
    else:
        ph = slide.shapes.add_shape(1, img_left, img_top, img_w, Inches(4.5))
        ph.fill.solid()
        ph.fill.fore_color.rgb = RGBColor(0xEE, 0xF2, 0xF6)
        ph.line.color.rgb = ACCENT
        tb = slide.shapes.add_textbox(img_left, img_top + Inches(1.8), img_w, Inches(1))
        tb.text_frame.paragraphs[0].text = f"（请运行 python tools/gen_screenshots.py 生成截图）"
    cap = slide.shapes.add_textbox(Inches(0.8), Inches(6.1), prs.slide_width - Inches(1.6), Inches(0.6))
    p = cap.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = caption
    r.font.size = Pt(14)
    r.font.color.rgb = GRAY
    r.font.name = "Microsoft YaHei"
    _notes(slide, notes)


def add_example_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _title_bar(slide, prs, "分析结果示例（真实输出节选）")
    examples = [
        ("客户异议", "【切换风险】客户原话：「切换影响营业」→ 建议：先 1 家店试点，老店照常开单"),
        ("待改进", "❌「保证三天上线」→ 改为「通常 3-5 个工作日，视门店数量而定」"),
        ("合规检查", "⚠️ 过度承诺风险：避免绝对化表述"),
        ("下一步", "48h 内发《平滑切换方案》+ 约下次 15 分钟演示"),
    ]
    y = Inches(1.5)
    for label, text in examples:
        box = slide.shapes.add_shape(1, Inches(0.8), y, prs.slide_width - Inches(1.6), Inches(1.15))
        box.fill.solid()
        colors = {"客户异议": RGBColor(0xFF, 0xF8, 0xE1), "待改进": RGBColor(0xFF, 0xF4, 0xF4),
                  "合规检查": RGBColor(0xFF, 0xF8, 0xE1), "下一步": RGBColor(0xE8, 0xF3, 0xFC)}
        box.fill.fore_color.rgb = colors.get(label, RGBColor(0xF5, 0xF9, 0xFD))
        box.line.fill.background()
        tb = slide.shapes.add_textbox(Inches(1.0), y + Inches(0.1), prs.slide_width - Inches(2.0), Inches(1.0))
        tf = tb.text_frame
        p1 = tf.paragraphs[0]
        r1 = p1.add_run()
        r1.text = label
        r1.font.bold = True
        r1.font.size = Pt(15)
        r1.font.color.rgb = PRIMARY
        r1.font.name = "Microsoft YaHei"
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = text
        r2.font.size = Pt(13)
        r2.font.color.rgb = DARK
        r2.font.name = "Microsoft YaHei"
        y += Inches(1.25)
    _notes(slide, "这是工具真实输出的节选。它不只是列问题，还会引用原话、分类异议、给出具体改法。"
                 "比如自动揪出「保证三天上线」这种过度承诺，并告诉你该怎么改。")


def build(out_path: str):
    # 尝试生成截图
    try:
        import tools.gen_screenshots as gs  # noqa
        gs.main()
    except Exception:
        pass

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_cover(prs)

    add_bullets(prs, "痛点：复盘是销售的「隐形时间黑洞」", [
        ("打完一通电话，手动复盘、记笔记要 15-20 分钟", 0),
        ("容易遗漏：客户异议没接住、过度承诺没发现", 0),
        ("经验难沉淀：好的话术没人帮你提炼", 0),
        ("一天打 20 通电话，光复盘就占掉大半天", 0),
    ], "我们赚钱靠打电话，但复盘不做又不行。手动复盘又慢又容易漏，好的话术也沉淀不下来。")

    add_bullets(prs, "我做了什么：通话复盘助手", [
        ("用 Cursor 开发了一个网站，上传录音转写即可分析", 0),
        ("自动输出 10 个维度的复盘表：", 0),
        ("需求 / 预算 / 异议（含应对建议）/ 做得好的 / 待改进（含改法）", 1),
        ("遗漏动作 / 下一步 / 合规检查 / 金句 / 评分", 1),
        ("已部署上线，打开链接就能用，不用装软件", 0),
    ], "我做了一个通话复盘助手，上传脱敏后的转写文件，几秒出结果。"
                 "它从十个维度帮你复盘，而且待改进不只是说问题，还给你具体改法。")

    add_bullets(prs, "我的 AI 学习路径（用 Cursor）", [
        ("用 Cursor 把想法变成能用的工具（不需要会编程）", 0),
        ("学会整理需求：复盘要哪些维度、输出什么格式", 0),
        ("部署成网站：免费平台一键上线，分享给同事", 0),
    ], "整个过程我用 Cursor 完成。不懂代码也能做，关键是把需求想清楚：复盘要分析什么、输出什么。")

    add_image_slide(prs, "怎么用 · 第 1 步：打开网站，上传文件",
                    os.path.join(SHOT_DIR, "01_上传页面.png"),
                    "打开你部署好的网址 → 拖入脱敏后的 Excel/CSV/txt → 点「开始分析」",
                    "现场演示第一步：打开网站，上传一份脱敏的转写文件。支持 Excel，就是你们 CRM 导出的那种格式。")

    add_image_slide(prs, "怎么用 · 第 2 步：查看复盘结果",
                    os.path.join(SHOT_DIR, "02_分析结果.png"),
                    "几秒后自动生成结构化复盘，可下载 Markdown 保存",
                    "点分析后几秒出结果。有评分、客户画像、异议分类、待改进（带改法）、合规检查、金句。")

    add_example_slide(prs)

    add_bullets(prs, "效果对比", [
        ("复盘时间：20 分钟 → 2 分钟", 0),
        ("自动发现过度承诺（如「保证三天上线」）", 0),
        ("异议逐条分类 + 给出建议应对话术", 0),
        ("可复用金句自动沉淀，团队共享", 0),
    ], "算笔账：以前复盘一通二十分钟，现在两分钟。而且机器不会漏掉过度承诺这种坑。")

    add_bullets(prs, "数据安全（分享必讲）", [
        ("上传前必须脱敏：去掉真实姓名、电话、公司名", 0),
        ("公开链接只传脱敏后的演示数据", 0),
        ("真实录音只在公司允许的环境使用", 0),
    ], "强调一点：数据安全。所有演示都是脱敏的，公开链接不要传真实客户数据。")

    add_bullets(prs, "总结 & 现场演示", [
        ("我做了一个工具：通话复盘助手", 0),
        ("网址：[填入你的 Streamlit 链接]", 0),
        ("接下来现场演示：上传一份文件，看分析结果", 0),
        ("欢迎大家扫码/收藏链接，会后试用", 0),
    ], "总结：我用 Cursor 做了通话复盘助手，已经部署上线。"
                 "接下来现场演示。谢谢大家！（把最后一页的链接换成你自己的网址）")

    prs.save(out_path)
    print(f"[完成] 已生成：{out_path}（共 {len(prs.slides)} 页，含演讲稿备注）")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("-o", "--out", default=os.path.join(ROOT, "AI学习成果分享.pptx"))
    args = parser.parse_args()
    build(args.out)


if __name__ == "__main__":
    main()
